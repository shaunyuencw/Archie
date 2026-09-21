import hashlib,io,json,re,zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from ..domain.models import Source,Passage,uid
from ..domain.commands import DomainError

ROOT=Path(__file__).resolve().parents[4]
MAX_UPLOAD=10*1024*1024
CHUNK_CHARS=5000
BATCH_CHARS=24000

def split_passage(locator,text,heading='',page=None):
    # Bounded, explicit ranges even for a single oversized section.
    return [Passage(locator=locator+(f'/chars/{i}-{min(i+CHUNK_CHARS,len(text))}' if len(text)>CHUNK_CHARS else ''),text=text[i:i+CHUNK_CHARS],heading=heading,page=page) for i in range(0,len(text),CHUNK_CHARS) if text[i:i+CHUNK_CHARS].strip()]

def parse(data:bytes,name:str,kind='document',store=None)->Source:
    if len(data)>MAX_UPLOAD: raise DomainError('invalid_input','Upload exceeds 10 MiB',413)
    digest=hashlib.sha256(data).hexdigest()
    if store:
        with store.connect() as db: cached=db.execute('SELECT body FROM cache WHERE hash=?',(digest,)).fetchone()
        if cached: return Source.model_validate_json(cached[0])
    ext=Path(name).suffix.lower(); passages=[]; unsupported=[]
    try:
        if kind=='prompt':
            passages=split_passage('prompt/1',data.decode('utf-8'))
        elif ext=='.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist())>50*1024*1024: raise DomainError('invalid_input','Expanded document is too large')
                if any('vbaProject' in n for n in archive.namelist()): raise DomainError('unsupported_document','Macro-enabled content is not supported')
            doc=Document(io.BytesIO(data)); heading=''
            for i,p in enumerate(doc.paragraphs,1):
                if p.style.name.startswith('Heading'): heading=p.text
                passages+=split_passage(f'paragraph/{i}',p.text,heading)
            for t,table in enumerate(doc.tables,1):
                for r,row in enumerate(table.rows,1):
                    for c,cell in enumerate(row.cells,1): passages+=split_passage(f'table/{t}/row/{r}/cell/{c}',cell.text,f'Table {t}')
        elif ext=='.pdf':
            reader=PdfReader(io.BytesIO(data))
            if reader.is_encrypted: raise DomainError('unsupported_document','Encrypted PDF is unsupported')
            for i,page in enumerate(reader.pages,1):
                text=page.extract_text() or ''
                if len(text.strip())<10: unsupported.append(i)
                else: passages+=split_passage(f'page/{i}',text,page=i)
        else: raise DomainError('unsupported_document','Only DOCX and text-based PDF are supported')
    except DomainError: raise
    except Exception as e: raise DomainError('unsupported_document','Document could not be parsed safely') from e
    manifest=json.loads((ROOT/'fixtures/projects/manifest.json').read_text())
    demo_manifest=ROOT/'fixtures/demos/manifest.json'
    if demo_manifest.exists(): manifest+=json.loads(demo_manifest.read_text())
    entry=next((m for m in manifest if m['sha256']==digest),None)
    src=Source(id=uid(),name=Path(name).name,kind=kind,origin='prompt' if kind=='prompt' else 'bundled_demo' if entry else 'upload',sha256=digest,canonical_id=entry['source_id'] if entry else digest,version=entry['version'] if entry else '1.0',variants=[digest],passages=passages,unprocessed=[p.locator for p in passages],unsupported_pages=unsupported)
    if store:
        with store.connect() as db: db.execute('INSERT OR REPLACE INTO cache VALUES(?,?)',(digest,src.model_dump_json()))
    return src

def batch(source:Source):
    selected=[]; chars=0
    for p in source.passages:
        if p.locator not in source.unprocessed: continue
        if len(selected)>=60 or chars+len(p.text)>BATCH_CHARS: break
        selected.append(p); chars+=len(p.text)
    return selected
