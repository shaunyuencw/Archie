import base64,csv,io,json,re
from html import escape
from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml.ns import qn
from ..domain.models import Project,uid
from ..domain.views import view_graph,narrative,initialise_views
ROOT=Path(__file__).resolve().parents[4]

def csv_text(p):
    stream=io.StringIO(newline='');writer=csv.writer(stream)
    fields=['id','source','target','purpose','data_direction','initiator','protocol','port','enforcement','delivery','semantic_revision']
    def safe(v):
        text='' if v is None else str(v)
        return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text
    writer.writerow(fields)
    for i in p.interfaces:
        values=i.model_dump();values['enforcement']=';'.join(i.enforcement);values['semantic_revision']=p.revision
        writer.writerow([safe(values.get(f)) for f in fields])
    return stream.getvalue()

def _wrap_label(text, width, font_size=13):
    """Keep exported text inside its presentation width, including long IDs."""
    import textwrap
    return [line for paragraph in text.splitlines() or ['']
            for line in textwrap.wrap(paragraph, max(8, int(width / (font_size * .56))))] or ['']


def _anchor(n, side):
    # Mirrors Canvas.tsx: side ports align with the icon, bottom exits below text.
    x, y, w, h = n['x'], n['y'], n['width'], n['height']
    return {'left': (x, y + 27), 'right': (x + w, y + 27),
            'top': (x + w / 2, y), 'bottom': (x + w / 2, y + h)}[side]


def _orthogonal_segment(start, end, source_side, target_side, offset=24):
    """Simple handle-aware elbows with the canvas's 24px port clearance."""
    vectors = {'left': (-1, 0), 'right': (1, 0), 'top': (0, -1), 'bottom': (0, 1)}
    u, v = vectors[source_side], vectors[target_side]
    a = (start[0] + u[0] * offset, start[1] + u[1] * offset)
    b = (end[0] + v[0] * offset, end[1] + v[1] * offset)
    if u[0] and v[0]:
        if (b[0] - a[0]) * u[0] >= 0:
            middle = [(sum((a[0], b[0])) / 2, a[1]), (sum((a[0], b[0])) / 2, b[1])]
        else:
            middle = [(a[0], (a[1] + b[1]) / 2), (b[0], (a[1] + b[1]) / 2)]
    elif u[1] and v[1]:
        middle = [(a[0], (a[1] + b[1]) / 2), (b[0], (a[1] + b[1]) / 2)]
    else:
        middle = [(b[0], a[1])] if u[0] else [(a[0], b[1])]
    points = [start, a, *middle, b, end]
    return [point for i, point in enumerate(points) if not i or point != points[i - 1]]


def _edge_points(a, b, route):
    source_side = route.source_handle if route else 'right'
    target_side = route.target_handle if route else 'left'
    start, end = _anchor(a, source_side), _anchor(b, target_side)
    waypoints = [(point.x, point.y) for point in route.points] if route else []
    if route and route.style == 'straight':
        return [start, *waypoints, end]
    if not waypoints:
        return _orthogonal_segment(start, end, source_side, target_side)
    points = [start]
    anchors = [start, *waypoints, end]
    for i, (first, last) in enumerate(zip(anchors, anchors[1:])):
        # Segment drags persist adjacent bend vertices, not extra midpoint hints.
        if first[0] == last[0] or first[1] == last[1]:
            if last != points[-1]:
                points.append(last)
            continue
        dx, dy = last[0] - first[0], last[1] - first[1]
        outgoing = ('right' if dx >= 0 else 'left') if abs(dx) >= abs(dy) else ('bottom' if dy >= 0 else 'top')
        opposite = {'left': 'right', 'right': 'left', 'top': 'bottom', 'bottom': 'top'}
        points.extend(_orthogonal_segment(first, last, source_side if i == 0 else outgoing,
                      target_side if i == len(anchors) - 2 else opposite[outgoing], offset=0)[1:])
    return points


def _edge_label_position(points, nodes, width, height, occupied):
    def clear(x, y):
        box = (x - width / 2 - 6, y - height / 2 - 6, x + width / 2 + 6, y + height / 2 + 6)
        obstacles = [(n['x'], n['y'], n['x'] + n['width'], n['y'] + n['height'])
                     for n in nodes.values() if n['role'] != 'zone'] + occupied
        return not any(box[0] < r[2] and box[2] > r[0] and box[1] < r[3] and box[3] > r[1] for r in obstacles)
    candidates = []
    for a, b in zip(points, points[1:]):
        length = abs(b[0] - a[0]) + abs(b[1] - a[1])
        for fraction in [.5, .25, .75]:
            x, y = a[0] + (b[0] - a[0]) * fraction, a[1] + (b[1] - a[1]) * fraction
            candidates.append((clear(x, y), length, -abs(fraction - .5), x, y))
    _, _, _, x, y = max(candidates)
    occupied.append((x - width / 2, y - height / 2, x + width / 2, y + height / 2))
    return x, y


def svg(p, kind):
    graph = view_graph(p, kind)
    nodes = {n['id']: dict(n) for n in graph['nodes']}
    for n in nodes.values():
        if n['parentId'] in nodes:
            parent = nodes[n['parentId']]
            n['x'] += parent['x']; n['y'] += parent['y']
    rendered_edges = [(e, _edge_points(nodes[e['source']], nodes[e['target']], e['route']))
                      for e in graph['edges'] if e['source'] in nodes and e['target'] in nodes]
    labels, occupied = {}, []
    for e, points in rendered_edges:
        lines = _wrap_label(e['label'], 260, 10)
        label_width = min(280, max(len(line) for line in lines) * 5.6 + 12)
        x, y = _edge_label_position(points, nodes, label_width, len(lines) * 14 + 6, occupied)
        labels[e['id']] = (x, y, label_width, lines)
    route_points = [point for _, points in rendered_edges for point in points]
    x0 = min([n['x'] for n in nodes.values()] + [x for x, _ in route_points] + [r[0] for r in occupied] + [0]) - 40
    y0 = min([n['y'] for n in nodes.values()] + [y for _, y in route_points] + [r[1] for r in occupied] + [0]) - 80
    width = max([n['x'] + n['width'] for n in nodes.values()] + [x for x, _ in route_points] + [r[2] for r in occupied] + [800]) - x0 + 40
    height = max([n['y'] + max(n['height'], 70 + len(_wrap_label(n['label'], n['width'] - 18)) * 17)
                  for n in nodes.values()] + [y for _, y in route_points] + [r[3] for r in occupied] + [500]) - y0 + 40
    label = 'SYNTHETIC — DEMONSTRATION ONLY' if p.synthetic else 'User-supplied project'
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{x0} {y0} {width} {height}" role="img">'
             f'<title>{escape(p.name)} — {kind} revision {p.revision}</title>'
             f'<metadata>Image export; not native Visio. {label}</metadata>'
             '<defs><marker id="flow-arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto-start-reverse" markerUnits="strokeWidth"><path d="M 0 0 L 8 4 L 0 8 z" fill="#587d93"/></marker></defs>'
             f'<rect x="{x0}" y="{y0}" width="{width}" height="{height}" fill="#f7fafc"/>'
             f'<text x="{x0 + 24}" y="{y0 + 30}" font-family="Arial" font-size="14" fill="#173c58">{escape(label)} · {kind} · revision {p.revision}</text>']
    from ..domain.catalogue import catalogue
    allowed = {a['id'] for a in catalogue()}

    def draw_node(n, zone=False):
        x, y, w, h = n['x'], n['y'], n['width'], n['height']
        parts.append(f'<g id="{escape(n["id"], quote=True)}" data-object-ids="{escape(json.dumps(n["object_ids"]), quote=True)}" class="{"zone" if zone else "equipment"}"><title>{escape(n["label"])}</title>')
        if zone:
            parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="8" fill="#eef5f7" fill-opacity=".2" stroke="#b29e76" stroke-width="2" stroke-dasharray="7 5"/>')
        elif n['asset_id'] in allowed:
            data = base64.b64encode((ROOT / 'fixtures/assets' / f'{n["asset_id"]}.svg').read_bytes()).decode()
            ix=x+(w-43)/2; iy=y+5
            quantity=n.get('quantity')
            for offset in reversed(range(1,min(quantity or 1,3))):
                parts.append(f'<image class="stack-copy" x="{ix+offset*6}" y="{iy-offset*4}" width="43" height="43" opacity=".35" href="data:image/svg+xml;base64,{data}"/>')
            parts.append(f'<image x="{ix}" y="{iy}" width="43" height="43" href="data:image/svg+xml;base64,{data}"/>')
            if quantity and quantity>1:
                mode={'unknown':'arrangement not decided','active_passive':'active / standby','active_active':'active / active'}[n['redundancy_mode']]
                parts.append(f'<text class="instance-count" x="{ix+47}" y="{iy+41}" font-family="Arial" font-size="11" fill="#187361"><title>{quantity} instances; {mode}</title>×{quantity}</text>')
            marker='AWS' if n['asset_id'].startswith('aws-') else 'VM' if n.get('form_factor')=='virtual' else None
            if marker:
                parts.append(f'<text x="{ix-8}" y="{iy+44}" font-family="Arial" font-size="8" fill="#526e86">{marker}</text>')
            if n.get('hosted_controls'):
                badge=base64.b64encode((ROOT/'fixtures/assets/virtual-firewall.svg').read_bytes()).decode()
                names=', '.join(c['name'] for c in n['hosted_controls'])
                parts.append(f'<g class="hosted-control"><title>{escape(names)} hosted here</title><rect x="{ix+44}" y="{iy-1}" width="27" height="27" rx="4" fill="#eef7f6" stroke="#b4d4cf"/><image x="{ix+47}" y="{iy+2}" width="21" height="21" href="data:image/svg+xml;base64,{badge}"/></g>')
        lines = _wrap_label(n['label'], w - 26 if zone else w - 18)
        for i, line in enumerate(lines):
            parts.append(f'<text x="{x + 13 if zone else x + w / 2}" y="{y + (27 if zone else 65) + i * 17}" font-family="Arial" font-size="{14 if zone else 13}" font-weight="600" text-anchor="{"start" if zone else "middle"}" fill="{"#655c45" if zone else "#1d3545"}">{escape(line)}</text>')
        parts.append(f'<text x="{x + 13 if zone else x + w / 2}" y="{y + (27 if zone else 65) + len(lines) * 17}" font-family="Arial" font-size="9" text-anchor="{"start" if zone else "middle"}" fill="#6c8091">{escape("SECURITY ZONE" if zone else n["role"])}</text></g>')

    for n in nodes.values():
        if n['role'] == 'zone': draw_node(n, True)
    for e, points in rendered_edges:
        path = 'M ' + ' L '.join(f'{x} {y}' for x, y in points)
        color = '#2782d4' if e['label'].startswith('video') else '#2b8b72' if e['label'].startswith('events') else '#b9783a' if e['label'].startswith('management') else '#718096'
        markers = (' marker-end="url(#flow-arrow)"' if e['data_direction'] in ['source_to_target', 'bidirectional'] else '') + (' marker-start="url(#flow-arrow)"' if e['data_direction'] in ['target_to_source', 'bidirectional'] else '')
        x, y, label_width, lines = labels[e['id']]
        parts.append(f'<g id="{escape(e["id"], quote=True)}" data-interface-ids="{escape(json.dumps(e["object_ids"]), quote=True)}"><path d="{path}" fill="none" stroke="{color}" stroke-width="1.8"{markers}/>'
                     f'<rect x="{x - label_width / 2}" y="{y - len(lines) * 7 - 3}" width="{label_width}" height="{len(lines) * 14 + 6}" rx="3" fill="#f6f9fb" fill-opacity=".96"/>')
        for i, line in enumerate(lines):
            parts.append(f'<text x="{x}" y="{y - len(lines) * 7 + 10 + i * 14}" font-family="Arial" font-size="10" text-anchor="middle" fill="#466173">{escape(line)}</text>')
        parts.append('</g>')
    for n in nodes.values():
        if n['role'] != 'zone': draw_node(n)
    return ''.join(parts) + '</svg>'

def docx_bytes(p):
    doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
    for name in ['Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.color.rgb=RGBColor(0,0,0)
        for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(22)
    for line in narrative(p).splitlines():
        if line.startswith('# '):doc.add_heading(line[2:],0)
        elif line.startswith('## '):doc.add_heading(line[3:],1)
        elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet').paragraph_format.keep_together=True
        elif line:doc.add_paragraph(line)
    stream=io.BytesIO();doc.save(stream);return stream.getvalue()

def pattern(p):
    """Whitelist reusable semantics. No addresses, arbitrary labels, evidence or answers."""
    out=Project(name='Synthetic reusable architecture pattern',policy_ids=[])
    from ..domain.catalogue import catalogue
    allowed={a['id'] for a in catalogue()}
    mapping={x.id:f'{group}-{i+1}' for group in ['systems','zones','components','deployments','interfaces','decisions'] for i,x in enumerate(getattr(p,group))}
    d=out.model_dump();roles={'video management','operator','configuration','management integration','analytics','audit','firewall','workstation','external command system'}
    d['systems']=[{'id':mapping[x.id],'name':f'System {i+1}','scope':x.scope} for i,x in enumerate(p.systems)]
    d['zones']=[{'id':mapping[x.id],'name':f'Zone {i+1}'} for i,x in enumerate(p.zones)]
    d['components']=[{'id':mapping[x.id],'name':x.asset_id if x.asset_id in allowed else 'generic role','role':x.role if x.role in roles else 'generic role','asset_id':x.asset_id if x.asset_id in allowed else 'generic-role-a','scope':x.scope,'system_id':mapping.get(x.system_id),'status':'proposed'} for x in p.components]
    d['deployments']=[{'id':mapping[x.id],'component_id':mapping[x.component_id],'zone_id':mapping.get(x.zone_id)} for x in p.deployments]
    d['interfaces']=[{'id':mapping[x.id],'source':mapping[x.source],'target':mapping[x.target],'purpose':x.purpose if x.purpose in ['video','events','management','integration'] else None,'data_direction':'unknown'} for x in p.interfaces]
    d['decisions']=[{'id':f'question-{i+1}','question':f'Confirm {x.field if x.field in ["initiator","delivery","purpose","value"] else "architecture decision"} for the new project','field':x.field if x.field in ['initiator','delivery','purpose','value'] else 'resolution','target_id':mapping.get(x.target_id),'options':['Not decided'],'answer':None,'state':'unknown'} for i,x in enumerate(p.decisions)]
    # Only the pattern's own fictional constraints; project-specific constraints never escape.
    d['constraints']=[{'id':'pattern-synthetic','key':'synthetic_pattern','value':True}]
    for kind,v in p.views.items():
        for ident,pos in v.placements.items():
            if ident in mapping:d['views'][kind]['placements'][mapping[ident]]=pos.model_copy(update={'label':None}).model_dump()
        for ident,route in v.routes.items():
            if ident in mapping:d['views'][kind]['routes'][mapping[ident]]=route.model_dump()
    return initialise_views(Project.model_validate(d))

def export(p,format,view='logical'):
    if view not in p.views:raise ValueError('Unknown view')
    if format=='json':return p.model_dump_json(indent=2).encode(),'application/json'
    if format=='svg':return svg(p,view).encode(),'image/svg+xml'
    if format=='md':return narrative(p).encode(),'text/markdown'
    if format=='csv':return csv_text(p).encode('utf-8-sig'),'text/csv'
    if format=='docx':return docx_bytes(p),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format=='pattern':return pattern(p).model_dump_json(indent=2).encode(),'application/json'
    raise ValueError('Unsupported export format')
