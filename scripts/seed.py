"""Deterministic fixture generator. No runtime providers or eval retrieval."""
import csv, hashlib, io, json, sys, zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
import yaml
from docx import Document
from docx.shared import Inches, Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from xml.sax.saxutils import escape
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from apps.api.app.domain.models import Project, Source, Passage, Claim

ROOT=Path(__file__).resolve().parents[1]
FIXED='2026-09-21T00:00:00+00:00'
LABEL='SYNTHETIC — DEMONSTRATION ONLY'
ASSETS=['server','gpu-server','application','analytics-engine','database','storage','firewall','switch','router','gateway','proxy','load-balancer','workstation','camera','identity-service','audit-service','event-broker','api-service','external-system','network-zone','system-boundary','site','generic-role-a','generic-role-b']
ICON_NAMES=['server','cpu','app-window','brain-circuit','database','hard-drive','shield','network','router','door-open','waypoints','git-fork','monitor','camera','key-round','scroll-text','radio-tower','plug','globe','square-dashed','group','building','box','circle-user']

def save(path,value):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')

def facts(pack):
    src=yaml.safe_load((ROOT/'fixtures/projects/source.yaml').read_text(encoding='utf-8'))
    d={k:deepcopy(src[k]) for k in ['systems','zones','components','interfaces','constraints']}
    if pack=='B':
        for i in d['interfaces']:
            i['initiator']=None
            if i['id']=='events': i['delivery']=None; i['protocol']=None
            if i['id']=='administration': i['purpose']=None; i['enforcement']=[]
        for c in d['constraints']:
            if c['key']=='availability_strategy': c['value']=None
    if pack=='C':
        d['components'].append(dict(id='cloud',name='External enrichment',role='internet dependency',asset_id='external-system',scope='external',system_id=None,zone_id=None,status='proposed',internet_hosted=True))
        d['interfaces'].append(dict(id='dependency',source='analytics',target='cloud',purpose='required enrichment',data_direction='bidirectional',initiator='analytics',protocol='HTTPS'))
        for c in d['constraints']:
            if c['key']=='resilience_required': c['value']=True
            if c['key']=='availability_strategy': c['value']='single_instance_no_recovery'
    return src,d

def statements(d,kind,pack):
    rows=[]
    for group in ['systems','zones','components','interfaces','constraints']:
        for item in d[group]:
            if pack=='C' and kind=='sow' and (item['id'] in ['cloud','dependency','availability']): continue
            if pack=='C' and kind=='spec' and item['id'] in ['offline','resilience']: continue
            # Human-readable lossless facts, not expected answers. Both render formats share these statements.
            fields='; '.join(f'{k}={json.dumps(v,ensure_ascii=False)}' for k,v in item.items() if k!='id')
            rows.append((item['id'],f'{group[:-1].capitalize()} {item["id"]}: {fields}'))
    return rows

def render_doc(path,title,summary,rows):
    doc=Document(); section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.65)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(5)
    doc.core_properties.created=doc.core_properties.modified=datetime(2026,9,21,tzinfo=timezone.utc)
    doc.core_properties.author='Architecture Workbench synthetic generator'
    doc.add_heading(title,0); doc.add_paragraph(LABEL); doc.add_paragraph('Version 1.0 | Fixture date 21 September 2026')
    doc.add_heading('Purpose and scope',1); doc.add_paragraph(summary)
    split=next(i for i,r in enumerate(rows) if r[1].startswith('Interface'))
    for n,part in enumerate([rows[:split],rows[split:]]):
        if n: doc.add_page_break()
        doc.add_heading('Systems and placement' if n==0 else 'Interfaces and decisions',1)
        table=doc.add_table(rows=1,cols=2); table.style='Light Shading Accent 1'
        table.rows[0].cells[0].text='ID'; table.rows[0].cells[1].text='Source requirement'
        table.columns[0].width=Inches(1); table.columns[1].width=Inches(5.8)
        for ident,line in part:
            cells=table.add_row().cells; cells[0].text=ident; cells[1].text=line
            for p in cells[1].paragraphs:
                for run in p.runs: run.font.size=Pt(8)
        doc.add_paragraph('Review guidance: null means not specified. Data direction is distinct from the session initiator. Icons do not prove enforcement or resilience. Equipment sizing needs a separate human decision.')
    buf=io.BytesIO(); doc.save(buf)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info=zipfile.ZipInfo(name,(2026,9,21,0,0,0)); info.compress_type=zipfile.ZIP_DEFLATED
            dst.writestr(info,src.read(name))

def render_pdf(path,title,summary,rows):
    styles=getSampleStyleSheet(); styles['BodyText'].fontSize=8; styles['BodyText'].leading=11
    story=[Paragraph(escape(title),styles['Title']),Paragraph(LABEL,styles['BodyText']),Paragraph('Version 1.0 | Fixture date 21 September 2026',styles['BodyText']),Spacer(1,12),Paragraph('Purpose and scope',styles['Heading1']),Paragraph(escape(summary),styles['BodyText']),Spacer(1,10)]
    split=next(i for i,r in enumerate(rows) if r[1].startswith('Interface'))
    for n,part in enumerate([rows[:split],rows[split:]]):
        if n: story.append(PageBreak())
        story.append(Paragraph('Systems and placement' if n==0 else 'Interfaces and decisions',styles['Heading1']))
        cells=[[Paragraph('ID',styles['BodyText']),Paragraph('Source requirement',styles['BodyText'])]]+[[Paragraph(escape(a),styles['BodyText']),Paragraph(escape(b),styles['BodyText'])] for a,b in part]
        table=Table(cells,colWidths=[75,440],repeatRows=1)
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dbe8f1')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f1f5f9')]),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.extend([table,Spacer(1,12),Paragraph('Review guidance: null means not specified. Do not invent addresses, ports, firewalls or server counts. C2 zoning remains unspecified outside VAP.',styles['BodyText'])])
    SimpleDocTemplate(str(path),pagesize=(612,792),leftMargin=48,rightMargin=48,topMargin=36,bottomMargin=36,invariant=1).build(story)

def reference(pack,d,rows):
    p=Project(id=f'reference-{pack}',name=f'{pack} reference — {LABEL}',created_at=FIXED,updated_at=FIXED)
    data=p.model_dump(); data.update({k:v for k,v in d.items() if k!='components'})
    data['components']=[]; data['deployments']=[]
    for c in d['components']:
        comp=dict(c); zone=comp.pop('zone_id',None); qty=comp.pop('quantity',None)
        data['components'].append(comp)
        data['deployments'].append(dict(id='deployment-'+c['id'],component_id=c['id'],zone_id=zone,quantity=qty))
    source=Source(id='reference-source',name=f'{pack} authored facts',kind='document',sha256='fixture',canonical_id=f'{pack}-reference',passages=[Passage(locator=f'row/{i+1}',text=line) for i,(_,line) in enumerate(rows)])
    data['sources']=[source.model_dump()]
    data['claims']=[Claim(id=f'claim-{ident}',source_id=source.id,source_version='1.0',locator=f'row/{i+1}',excerpt=line,target_id=ident,field='record',value=line,source_kind='document',review='confirmed').model_dump() for i,(ident,line) in enumerate(rows)]
    for c in data['components']+data['interfaces']: c['evidence']=['claim-'+c['id']]
    from apps.api.app.domain.views import initialise_views
    return initialise_views(Project.model_validate(data))

def main(out=None):
    target=Path(out or ROOT/'fixtures'); target.mkdir(parents=True,exist_ok=True); manifest=[]
    for pack in 'ABC':
        src,d=facts(pack); meta=src['packs'][pack]
        for kind in ['sow','spec']:
            folder=target/'projects'/pack; folder.mkdir(parents=True,exist_ok=True)
            rows=statements(d,kind,pack); title=f'{meta["title"]} {"statement of work" if kind=="sow" else "technical specification"}'
            (folder/f'{kind}.md').write_text(f'# {title}\n\n{LABEL}\n\n{meta["summary"]}\n\n'+'\n\n'.join(line for _,line in rows)+'\n',encoding='utf-8')
            for ext,render in [('docx',render_doc),('pdf',render_pdf)]:
                path=folder/f'{kind}.{ext}'; render(path,title,meta['summary'],rows)
                manifest.append({'path':str(path.relative_to(target)).replace('\\','/'),'source_id':f'{pack}-{kind}','version':'1.0','sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
        if pack in 'AB':
            p=reference(pack,d,statements(d,'spec',pack)); folder=target/'references'/pack
            save(folder/'project.json',p.model_dump()); save(folder/'views.json',{k:v.model_dump() for k,v in p.views.items()})
            (folder/'narrative.md').write_text(f'# {p.name}\n\nSemantic revision 0.\n\n'+'\n'.join(f'- {c.id}: {c.name}; {c.role}.' for c in p.components)+'\n\nUnknowns: '+('event delivery, initiators, administration and availability.' if pack=='B' else 'equipment sizing and ports are deliberately undecided.'),encoding='utf-8')
            with (folder/'interfaces.csv').open('w',newline='',encoding='utf-8') as f:
                w=csv.writer(f); w.writerow(['id','source','target','purpose','data_direction','initiator','protocol'])
                for i in p.interfaces: w.writerow([i.id,i.source,i.target,i.purpose,i.data_direction,i.initiator,i.protocol])
            save(folder/'expected-findings.json',{'label':LABEL,'DEMO-IF-01':'pass' if pack=='A' else 'insufficient_information','DEMO-ZON-01':'pass','DEMO-HA-01':'not_applicable'})
    save(target/'projects/manifest.json',manifest)
    policies=yaml.safe_load((ROOT/'fixtures/policies/source.yaml').read_text(encoding='utf-8'))
    for c in policies['clauses']: c['version']='1.0'; c['applicability']=c['tags']
    save(target/'policies/clauses.json',policies)
    catalogue=[]
    for asset,icon in zip(ASSETS,ICON_NAMES):
        source=ROOT/f'apps/web/node_modules/lucide-static/icons/{icon}.svg'
        folder=target/'assets'; folder.mkdir(exist_ok=True)
        (folder/f'{asset}.svg').write_bytes(source.read_bytes())
        catalogue.append({'id':asset,'type':asset,'icon':f'{asset}.svg','tags':asset.split('-'),'handles':['left','right','top','bottom'],'size':{'width':170,'height':85},'licence':'ISC / Lucide 0.468.0'})
    save(target/'assets/manifest.json',catalogue)
    (target/'assets/LICENCE.txt').write_bytes((ROOT/'apps/web/node_modules/lucide-static/LICENSE').read_bytes())
    names=['Prompt baseline','DOCX table extraction','PDF version handling','Unknowns and questions','Source contradiction','Policy and coverage','Mouse operations','Mouse then prompt','Version conflict','Views and narrative','Provider failures','Budget provider isolation','Injection variant','Capacity variant','Oversized document variant','Export template variant']
    expectations=['8 components, 4 interfaces, 4 zones; sizing undecided','Table locator resolves exact excerpt','Page evidence and format deduplication','Null fields and at most three questions; answers persist','Both dependency and availability sources retained','8 implemented plus 12 manual clauses','Manual operations have zero provider calls','Unrelated placement and routes unchanged','Stale semantic and presentation bases rejected','All derived outputs use accepted revision','Malformed, fabricated references and timeout cannot commit','Persistent atomic reservations and zero hidden fallback','Host tools and budgets cannot be changed by source text','50 components and 100 interfaces survive edit and reopen','Explicit processed and unprocessed ranges','Cross-export IDs; sanitised pattern; native Visio separate']
    for i,(name,expected) in enumerate(zip(names,expectations),1): save(target/f'evals/{"development" if i<=12 else "frozen"}/T{i:02}.json',{'id':f'T{i:02}','name':name,'expected':expected,'version':'1.0','authoring':'Agent-authored fixture; not an independent benchmark','runtime_retrieval':False})
    print(f'Generated six documents in two formats, 20 clauses, 24 assets, two references and 16 scenarios at {target}')

if __name__=='__main__': main(sys.argv[1] if len(sys.argv)>1 else None)
