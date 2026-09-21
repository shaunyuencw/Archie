"""Generate bounded synthetic demo packs from source.yaml without provider calls.

Existing A/B reference fixtures are intentionally untouched.
"""
import argparse,hashlib,io,json,sys,zipfile
from copy import deepcopy
from datetime import datetime,timezone
from pathlib import Path
from xml.sax.saxutils import escape
import yaml
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
from reportlab.lib.styles import getSampleStyleSheet
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from apps.api.app.domain.models import Project,Source,Claim,Placement,Route
from apps.api.app.domain.views import initialise_views,narrative,view_graph
from apps.api.app.ingest.parser import parse
from apps.api.app.exports.service import export

ROOT=Path(__file__).resolve().parents[1]
FIXED='2026-09-21T00:00:00+00:00'
LABEL='SYNTHETIC — DEMONSTRATION ONLY'


def save_json(path,value):
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def canonical_docx(document,path):
    document.core_properties.created=document.core_properties.modified=datetime(2026,9,21,tzinfo=timezone.utc)
    document.core_properties.author='ARCHIE synthetic demo generator'
    stream=io.BytesIO();document.save(stream)
    with zipfile.ZipFile(stream) as source,zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info=zipfile.ZipInfo(name,(2026,9,21,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED
            target.writestr(info,source.read(name))


def document_base():
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.75)
    doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    for name in ['Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.color.rgb=RGBColor(0,0,0)
        for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(22)
    doc.styles['Heading 1'].font.size=Pt(13)
    return doc


def render_spec(folder,scenario):
    title=scenario['title']+' technical specification'
    doc=document_base();doc.add_heading(title,0);doc.add_paragraph(LABEL)
    doc.add_paragraph('Version 1.0 · Fixed fixture date 21 September 2026')
    doc.add_paragraph(scenario['summary'])
    styles=getSampleStyleSheet();styles['BodyText'].fontSize=10;styles['BodyText'].leading=13
    styles['BodyText'].spaceAfter=8;styles['Title'].fontSize=21;styles['Title'].leading=25
    styles['Heading1'].fontSize=13;styles['Heading1'].leading=16
    story=[Paragraph(escape(title),styles['Title']),Paragraph(escape(LABEL),styles['BodyText']),Paragraph('Version 1.0 · Fixed fixture date 21 September 2026',styles['BodyText']),Paragraph(escape(scenario['summary']),styles['BodyText'])]
    markdown=[f'# {title}','',LABEL,'','Version 1.0 · Fixed fixture date 21 September 2026','',scenario['summary'],'']
    section=None
    for item in scenario['requirements']:
        if item['section']!=section:
            if scenario['featured'] and item['section']=='Interfaces':doc.add_page_break();story.append(PageBreak())
            section=item['section'];doc.add_heading(section,1);story.append(Paragraph(escape(section),styles['Heading1']));markdown.extend(['## '+section,''])
        text=f'[{item["id"]}] {item["text"]}'
        paragraph=doc.add_paragraph(text);paragraph.paragraph_format.keep_together=True
        story.append(Paragraph(escape(text),styles['BodyText']));markdown.extend([text,''])
    canonical_docx(doc,folder/'techspec.docx')
    SimpleDocTemplate(str(folder/'techspec.pdf'),pagesize=(612,792),leftMargin=48,rightMargin=48,topMargin=40,bottomMargin=40,invariant=1).build(story)
    (folder/'techspec.md').write_text('\n'.join(markdown),encoding='utf-8')


def records(scenario):
    groups={key:[] for key in ['systems','zones','components','deployments','interfaces','constraints','decisions']}
    refs={}
    for group in ['systems','zones','components','interfaces','constraints','decisions']:
        for original in scenario[group]:
            item=deepcopy(original);reference=item.pop('ref');refs[item['id']]=reference
            if group=='components':
                deployment={'id':'deployment-'+item['id'],'component_id':item['id']}
                for key in ['zone_id','quantity','redundancy_mode','host_component_id']:deployment[key]=item.pop(key,None)
                groups['deployments'].append(deployment);refs[deployment['id']]=reference
            groups[group].append(item)
    return groups,refs


def place_project(project,code):
    """Deliberate review layouts; points are presentation, never security evidence."""
    if code=='portal':
        zones={'client':(20,30,225,430),'z1':(265,30,420,430),'z2':(705,30,225,430)}
        nodes={'employee':(20,150),'service-desk':(20,290),'portal-web':(15,150),'portal-app':(220,150),'app-host':(15,290),'zone-firewall':(220,290),'request-db':(20,150)}
        routes={'employee-web':('top','top',[(132.5,110),(372.5,110)]),
                'desk-web':('right','left',[(255,347),(255,207)]),
                'web-app':('top','top',[(372.5,145),(577.5,145)]),
                'app-database':('top','top',[(577.5,110),(817.5,110)])}
    else:
        zones={'prod-client':(20,30,235,200),'prod-z1':(275,30,480,200),'prod-z2':(775,30,235,200),
               'test-client':(20,260,235,440),'test-z1':(275,260,480,440),'test-z2':(775,260,235,440),
               'aws-z1':(275,740,480,225),'aws-z2':(775,740,235,225)}
        nodes={'prod-operator':(25,70),'prod-api':(145,70),'prod-db':(25,70),'trial-robot':(25,70),'trial-tablet':(25,195),'trial-ap':(25,320),
               'robot-controller':(15,70),'trial-gateway':(280,195),'trial-host':(15,320),'trial-firewall':(280,320),'trial-data':(25,70),
               'aws-api':(15,65),'aws-function':(280,65),'aws-telemetry':(25,65)}
        routes={'production-access':('right','left',[]),'production-database':('right','left',[]),'trial-wireless':('left','left',[(30,357),(30,607)]),
                'robot-telemetry':('left','left',[(10,357),(10,240),(265,240),(265,357)]),
                'trial-operator-access':('right','bottom',[(382.5,482)]),'trial-database':('right','left',[]),
                'trial-snapshot':('top','right',[(647.5,415),(505,415),(505,357)]),
                'production-handoff':('right','bottom',[(765,482),(765,245),(512.5,245)]),
                'cloud-telemetry':('bottom','left',[(647.5,570),(765,570),(765,720),(265,720),(265,832)]),
                'telemetry-transform':('bottom','bottom',[(382.5,940),(647.5,940)]),
                'telemetry-store':('bottom','bottom',[(647.5,940),(892.5,940)])}
    for kind in ['logical','sv2']:
        view=project.views[kind]
        for ident,(x,y,w,h) in zones.items():view.placements[ident]=Placement(x=x,y=y,width=w,height=h)
        for ident,(x,y) in nodes.items():view.placements[ident]=Placement(x=x,y=y,width=185,height=115)
        for ident,(source,target,points) in routes.items():view.routes[ident]=Route(source_handle=source,target_handle=target,points=[{'x':x,'y':y} for x,y in points],locked=bool(points))
    summary=project.views['sv1']
    for i,system in enumerate(project.systems):summary.placements[system.id]=Placement(x=25+i*350,y=50,width=320,height=190)
    # Aggregate labels need a long clear rail rather than the narrow gap between groups.
    for i,edge in enumerate(view_graph(project,'sv1')['edges']):
        a,b=summary.placements[edge['source']],summary.placements[edge['target']]
        y=280+i*50
        summary.routes[edge['id']]=Route(source_handle='bottom',target_handle='bottom',points=[{'x':a.x+a.width/2,'y':y},{'x':b.x+b.width/2,'y':y}],locked=True)
    return initialise_views(project)


def build_reference(folder,scenario):
    groups,refs=records(scenario)
    source=parse((folder/'techspec.docx').read_bytes(),'techspec.docx')
    source.id=f'demo-{scenario["code"]}-source';source.canonical_id=f'demo-{scenario["code"]}-techspec';source.origin='bundled_demo'
    source.variants=[hashlib.sha256((folder/f'techspec.{ext}').read_bytes()).hexdigest() for ext in ['docx','pdf']]
    source.processed=[p.locator for p in source.passages];source.unprocessed=[]
    claims=[]
    for group,items in groups.items():
        for item in items:
            paragraph=next(p for p in source.passages if p.text.startswith('['+refs[item['id']]+']'))
            claim=Claim(id='claim-'+item['id'],source_id=source.id,source_version='1.0',locator=paragraph.locator,excerpt=paragraph.text,target_id=item['id'],field='record',value=deepcopy(item),source_kind='document',review='confirmed')
            claims.append(claim)
            if group in ['components','interfaces','constraints','decisions']:item['evidence']=[claim.id]
    project=Project(id='demo-'+scenario['code'],name=scenario['title'],created_at=FIXED,updated_at=FIXED,sources=[source],claims=claims,policy_ids=scenario.get('policy_ids'),**groups,
                    notes='Synthetic authored example. Confirmed claims mean confirmed within this fictional reference; they are not organisational approval. See Sources for the complete specification and open decisions.')
    project=place_project(project,scenario['code']);save_json(folder/'project.json',project.model_dump())
    (folder/'narrative.md').write_text(narrative(project)+'\n',encoding='utf-8')
    for kind in project.views:(folder/f'{kind}.svg').write_bytes(export(project,'svg',kind)[0])
    return project


def mock_input(folder,scenario):
    """Separate offline parser fixture, never presented as general NL understanding."""
    groups,_=records(scenario);deployments={d['component_id']:d for d in groups['deployments']}
    doc=document_base();doc.add_heading(scenario['title']+' authored mock input',0);doc.add_paragraph(LABEL)
    doc.add_paragraph('Deterministic authored records for the mock parser. Use techspec.docx or techspec.pdf with a live provider to test natural-language interpretation. No policy approval is implied.')
    lines=[]
    for group,record_name in [('systems','System'),('zones','Zone'),('components','Component'),('interfaces','Interface'),('constraints','Constraint')]:
        for original in groups[group]:
            item=deepcopy(original);ident=item.pop('id')
            if group=='components':item.update({k:v for k,v in deployments[ident].items() if k not in ['id','component_id']})
            text=f'{record_name} {ident}: '+'; '.join(f'{key}={json.dumps(value,ensure_ascii=False)}' for key,value in item.items())
            lines.append(text);doc.add_paragraph(text)
    canonical_docx(doc,folder/'mock-input.docx');(folder/'mock-input.txt').write_text(LABEL+'\n\n'+'\n\n'.join(lines)+'\n',encoding='utf-8')


def main(output_dir=None):
    source=yaml.safe_load((ROOT/'fixtures/demos/source.yaml').read_text(encoding='utf-8'))
    output=Path(output_dir or ROOT/'fixtures/demos');output.mkdir(parents=True,exist_ok=True);manifest=[]
    for scenario in source['scenarios']:
        folder=output/scenario['code'];folder.mkdir(parents=True,exist_ok=True)
        render_spec(folder,scenario);save_json(folder/'followups.json',{'synthetic':True,'scenario':scenario['title'],'provider_notes':'Mock supports authored key-value fixtures and narrow add-workstation/rename prompts only. Natural-language technical specs and all broader follow-ups require an explicitly selected live provider. Review every change set before accepting. Ordinary edits and exports use no model calls.','prompts':scenario['followups']})
        for ext in ['docx','pdf']:
            path=folder/f'techspec.{ext}';manifest.append({'path':f'{scenario["code"]}/techspec.{ext}','source_id':f'demo-{scenario["code"]}-techspec','version':'1.0','sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
        if scenario['featured']:build_reference(folder,scenario);mock_input(folder,scenario)
    save_json(output/'manifest.json',manifest)
    print(f'Generated four synthetic specifications and two layered references at {output}')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--output-dir',type=Path);main(parser.parse_args().output_dir)
